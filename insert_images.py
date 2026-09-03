"""Insert field notebook photos into the Word report at 📷 markers."""
import sys, io, os, re
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

DOCX_PATH = Path(r"C:\Users\asus\Desktop\1111 - 副本.docx")
OUT_PATH = Path(r"C:\Users\asus\Desktop\1111_带图片.docx")
PHOTO_DIR = Path(r"D:\geointern-agent\field_notebooks\周口店")
DAY1_DIR = PHOTO_DIR / "day1"

# Collect all available photos
all_photos = {}
for d in [PHOTO_DIR, DAY1_DIR]:
    for f in d.glob("*.jpg"):
        all_photos[f.name] = f
print(f"可用照片: {len(all_photos)} 张")

# Key photos to insert at specific locations (most representative per section)
# Mapping: paragraph text keyword -> list of photo filenames
SECTION_PHOTOS = {
    'day1/IMG20260722180324.jpg~0445.jpg（L06，9张）': [
        'IMG20260722180324.jpg', 'IMG20260722180432.jpg',
    ],
    '燧石条带露头照片': ['D0601燧石条带先后关系.jpg'],
    '叠层石露头照片': ['D0605包心菜状叠层石.jpg'],
    'IMG20260803162904.jpg、IMG20260803162908.jpg': [
        'IMG20260803162904.jpg', 'IMG20260803162908.jpg'
    ],
    'IMG20260803162759.jpg': ['IMG20260803162759.jpg'],
    'IMG20260803162744.jpg、IMG20260803162753.jpg（图7-1': [
        'IMG20260803162744.jpg', 'IMG20260803162753.jpg'
    ],
    'IMG20260803162802.jpg': ['IMG20260803162802.jpg'],
    'IMG20260803162811.jpg': ['IMG20260803162811.jpg'],
    'IMG20260803162813.jpg': ['IMG20260803162813.jpg'],
    'IMG20260803162816.jpg': ['IMG20260803162816.jpg'],
    'IMG20260803162820.jpg（L07冶里': ['IMG20260803162820.jpg'],
    'IMG20260803162944.jpg（L10马家沟组）': ['IMG20260803162944.jpg'],
    'IMG20260803163026.jpg（L11马家沟组）': ['IMG20260803163026.jpg'],
    'IMG20260803163029.jpg（L11': ['IMG20260803163029.jpg'],
    'IMG20260803162953.jpg（L10': ['IMG20260803162953.jpg'],
    'IMG20260803163035.jpg（L11': ['IMG20260803163035.jpg'],
    'IMG20260803163039.jpg（L11': ['IMG20260803163039.jpg'],
    'IMG20260803163043.jpg（L11': ['IMG20260803163043.jpg'],
    'IMG20260803162908.jpg（图9-2': ['IMG20260803162908.jpg'],
    'IMG20260803162831.jpg、IMG20260803162835.jpg': [
        'IMG20260803162831.jpg', 'IMG20260803162835.jpg'
    ],
    'IMG20260803162851.jpg': ['IMG20260803162851.jpg'],
    'IMG20260803162841.jpg、IMG20260803162853.jpg（图8-2': [
        'IMG20260803162841.jpg', 'IMG20260803162853.jpg'
    ],
    'IMG20260803162856.jpg（图8-3': ['IMG20260803162856.jpg'],
    'IMG20260803162858.jpg、IMG20260803162902.jpg（图8-4': [
        'IMG20260803162858.jpg', 'IMG20260803162902.jpg'
    ],
    'IMG20260803162913.jpg（羊屎沟': ['IMG20260803162913.jpg'],
    'IMG20260803162926.jpg（羊屎沟片岩': ['IMG20260803162926.jpg'],
    'IMG20260803162937.jpg、IMG20260803162942.jpg（L09路线小结': [
        'IMG20260803162937.jpg', 'IMG20260803162942.jpg'
    ],
    'IMG20260803162923.jpg（图9-3': ['IMG20260803162923.jpg'],
    'IMG20260803163048.jpg（L11路线小结': ['IMG20260803163048.jpg'],
    'IMG20260803163053.jpg~3118.jpg（L12）': [
        'IMG20260803163053.jpg', 'IMG20260803163104.jpg'
    ],
    # Section header broad references
    'day1/IMG20260722180324.jpg~0445.jpg（L06，9张）、IMG20260803162703': [
        'IMG20260722180324.jpg',
    ],
    'day1/IMG20260722180324.jpg~0445.jpg': [
        'IMG20260722180324.jpg',
    ],
    'day1/IMG20260722180432.jpg': [
        'IMG20260722180432.jpg',
    ],
    'IMG20260803162744.jpg~2829.jpg（L07全部）': [
        'IMG20260803162744.jpg',
    ],
    'IMG20260803162831.jpg~2902.jpg（L08全部）、IMG20260803162904.jpg~2942.jpg（L09）': [
        'IMG20260803162831.jpg', 'IMG20260803162904.jpg',
    ],
    'IMG20260803162904.jpg~2942.jpg（L09全部）': [
        'IMG20260803162904.jpg',
    ],
    'IMG20260803162753.jpg（L07景儿峪组-府君山组平卧褶皱 图7-1）': [
        'IMG20260803162753.jpg',
    ],
}

def compress_photo(photo_path):
    """Resize photo for docx insertion. Returns BytesIO."""
    img = Image.open(photo_path)
    # Target: fit within 1200x900
    max_w, max_h = 1200, 900
    ratio = min(max_w / img.size[0], max_h / img.size[1], 1.0)
    if ratio < 1:
        img = img.resize((int(img.size[0]*ratio), int(img.size[1]*ratio)))
    if img.mode == 'RGBA':
        img = img.convert('RGB')
    buf = io.BytesIO()
    img.save(buf, 'JPEG', quality=65)
    buf.seek(0)
    return buf

def resolve_photo(name):
    """Find a photo file by name (fuzzy match, handles day1/ prefix)."""
    # Strip day1/ prefix
    clean_name = name.replace('day1/', '')
    # Exact match
    if clean_name in all_photos:
        return all_photos[clean_name]
    # Partial match
    for k, v in all_photos.items():
        k_core = k.replace('.jpg','').lower()
        n_core = clean_name.replace('.jpg','').lower()
        if n_core in k_core or k_core in n_core:
            return v
    return None

def find_photos_for_range(start_name, end_name):
    """Find all photos alphabetically between start and end."""
    result = []
    for name in sorted(all_photos.keys()):
        if start_name <= name <= end_name:
            result.append(name)
    return result

# Read docx
doc = Document(DOCX_PATH)

# Collect all 📷 paragraphs
ref_paragraphs = []
for i, p in enumerate(doc.paragraphs):
    if '📷' in p.text:
        ref_paragraphs.append((i, p))

print(f"找到 {len(ref_paragraphs)} 处图片引用")

# Process: for each 📷 paragraph, find matching photos and insert
inserted_count = 0
for para_idx, para in ref_paragraphs:
    text = para.text
    photos_to_insert = []

    # Check each key against this paragraph
    for keyword, photo_names in SECTION_PHOTOS.items():
        if keyword in text:
            for pn in photo_names:
                resolved = resolve_photo(pn)
                if resolved and resolved not in photos_to_insert:
                    photos_to_insert.append(resolved)
            break

    if not photos_to_insert:
        continue

    # Clear the reference paragraph and insert images
    # First, clear existing text
    for run in para.runs:
        run.text = ''

    # Keep first run, clear others
    if para.runs:
        para.runs[0].text = ''
    # Remove extra runs
    for run in para.runs[1:]:
        run._element.getparent().remove(run._element)

    # Add images as inline shapes
    for idx, photo_path in enumerate(photos_to_insert[:2]):  # Max 2 per ref
        try:
            img_data = compress_photo(photo_path)
            # Calculate image width in cm (proportional to compressed image)
            img = Image.open(img_data)
            img_data.seek(0)
            w_cm = min(15, img.size[0] / 100)  # ~15cm max width
            h_cm = w_cm * img.size[1] / img.size[0]

            run = para.add_run()
            run.add_picture(img_data, width=Cm(w_cm), height=Cm(h_cm))
            if idx < len(photos_to_insert[:2]) - 1:
                para.add_run('  ')
            inserted_count += 1
        except Exception as e:
            print(f"  ERROR: {photo_path.name}: {e}")

    # Add small caption below images
    caption_run = para.add_run(f'\n[野簿: {text.replace("> 📷 野簿：", "").strip()[:80]}]')
    caption_run.font.size = Pt(8)
    caption_run.font.color.rgb = RGBColor(150, 150, 150)
    caption_run.font.italic = True

print(f"插入图片: {inserted_count} 张")

# Save
doc.save(str(OUT_PATH))
size_mb = OUT_PATH.stat().st_size / (1024*1024)
print(f"输出: {OUT_PATH}")
print(f"大小: {size_mb:.1f} MB")
